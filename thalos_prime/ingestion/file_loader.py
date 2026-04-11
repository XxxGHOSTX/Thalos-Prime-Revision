from __future__ import annotations

import hashlib
import io
import logging
from pathlib import Path

from thalos_prime.config import Settings
from thalos_prime.ingestion.models import IngestionResult
from thalos_prime.storage.db import RawDocument, get_session, init_db

logger = logging.getLogger(__name__)


class FileLoader:
    def __init__(self, settings: Settings):
        self.settings = settings

    def _store_raw(self, text: str, url: str, source_type: str) -> IngestionResult:
        init_db(self.settings.database_url)
        content_hash = hashlib.sha256(text.encode("utf-8")).hexdigest()
        with get_session(self.settings.database_url) as session:
            from sqlalchemy import select

            existing = session.execute(
                select(RawDocument).where(RawDocument.content_hash == content_hash)
            ).scalar_one_or_none()
            if existing:
                return IngestionResult(
                    doc_id=existing.id,
                    url=url,
                    status="duplicate",
                    word_count=len(text.split()),
                    message="Duplicate content hash",
                )
            doc = RawDocument(
                url=url,
                content_hash=content_hash,
                raw_text=text,
                source_type=source_type,
                robots_allowed=True,
                status="raw",
            )
            session.add(doc)
            session.flush()
            doc_id = doc.id

        return IngestionResult(
            doc_id=doc_id,
            url=url,
            status="ingested",
            word_count=len(text.split()),
        )

    def load_pdf(self, path: Path) -> IngestionResult:
        try:
            from pdfminer.high_level import extract_text as pdf_extract_text
            text = pdf_extract_text(str(path))
        except Exception as e:
            logger.error("PDF load failed: %s", e)
            return IngestionResult(doc_id=None, url=str(path), status="error", message=str(e))
        return self._store_raw(text, str(path), "pdf")

    def load_docx(self, path: Path) -> IngestionResult:
        try:
            from docx import Document
            doc = Document(str(path))
            text = "\n".join(para.text for para in doc.paragraphs)
        except Exception as e:
            logger.error("DOCX load failed: %s", e)
            return IngestionResult(doc_id=None, url=str(path), status="error", message=str(e))
        return self._store_raw(text, str(path), "docx")

    def load_text(self, path: Path) -> IngestionResult:
        try:
            import chardet
            raw_bytes = path.read_bytes()
            detected = chardet.detect(raw_bytes)
            encoding = detected.get("encoding") or "utf-8"
            text = raw_bytes.decode(encoding, errors="replace")
        except Exception as e:
            logger.error("Text load failed: %s", e)
            return IngestionResult(doc_id=None, url=str(path), status="error", message=str(e))
        return self._store_raw(text, str(path), "text")

    def load(self, path: Path) -> IngestionResult:
        path = Path(path)
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self.load_pdf(path)
        elif suffix in (".docx", ".doc"):
            return self.load_docx(path)
        else:
            return self.load_text(path)
